from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from hashlib import sha256
import csv
import io
import json
from pathlib import Path
import re

from openpyxl import load_workbook


HEADING_PATTERN = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
WORD_PATTERN = re.compile(r"\S+")
STRUCTURED_EXTENSIONS = {".csv", ".json", ".xlsx"}
PROGRESS_BAR_WIDTH = 40


def _display_progress(label: str, completed: int, total: int) -> None:
    if total == 0:
        print(f"{label} [no files]", flush=True)
        return

    filled_width = round(PROGRESS_BAR_WIDTH * completed / total)
    bar = "█" * filled_width + "░" * (PROGRESS_BAR_WIDTH - filled_width)
    line_ending = "\n" if completed == total else ""
    print(f"\r{label} [{bar}]", end=line_ending, flush=True)


@dataclass(frozen=True, slots=True)
class DocumentChunk:
    chunk_id: str
    content: str
    source_file: str
    source_type: str
    section: str
    chunk_index: int


def _split_markdown_sections(content: str) -> list[tuple[str, str]]:
    sections = []
    heading_stack: list[str] = []
    current_section = ""
    current_lines: list[str] = []

    def save_section() -> None:
        body = "\n".join(current_lines).strip()
        if body:
            sections.append((current_section, body))

    for line in content.splitlines():
        heading = HEADING_PATTERN.match(line)

        if not heading:
            current_lines.append(line)
            continue

        save_section()
        current_lines = []

        level = len(heading.group(1))
        title = heading.group(2).strip()
        heading_stack = heading_stack[: level - 1]
        heading_stack.append(title)
        current_section = " > ".join(heading_stack)

    save_section()
    return sections


def _split_by_words(
    content: str,
    chunk_size: int,
    chunk_overlap: int,
) -> list[str]:
    words = list(WORD_PATTERN.finditer(content))

    if not words:
        return []

    chunks = []
    step = chunk_size - chunk_overlap

    for start in range(0, len(words), step):
        end = min(start + chunk_size, len(words))
        character_start = words[start].start()
        character_end = words[end - 1].end()
        chunks.append(content[character_start:character_end].strip())

        if end == len(words):
            break

    return chunks


def _chunk_file(
    file_path: Path,
    source_directory: Path,
    chunk_size: int,
    chunk_overlap: int,
) -> list[DocumentChunk]:
    content = file_path.read_text(encoding="utf-8")
    sections = _split_markdown_sections(content)

    if not sections and content.strip():
        sections = [("", content.strip())]

    source_file = file_path.relative_to(source_directory).as_posix()
    chunks = []

    for section, section_content in sections:
        for piece in _split_by_words(section_content, chunk_size, chunk_overlap):
            chunk_index = len(chunks)
            chunk_content = (
                f"Section: {section}\n\n{piece}" if section else piece
            )
            chunk_id = sha256(
                f"{source_file}:{chunk_index}:{chunk_content}".encode("utf-8")
            ).hexdigest()[:16]

            chunks.append(
                DocumentChunk(
                    chunk_id=chunk_id,
                    content=chunk_content,
                    source_file=source_file,
                    source_type="markdown",
                    section=section,
                    chunk_index=chunk_index,
                )
            )

    return chunks


def chunk_documents(
    input_directory: str | Path,
    chunk_size: int,
    chunk_overlap: int,
    batch_size: int,
    max_workers: int,
) -> list[DocumentChunk]:
    """Chunk Markdown documents in parallel batches with console progress."""
    source_directory = Path(input_directory)

    if not source_directory.is_dir():
        raise NotADirectoryError(
            f"Temporary data directory not found: {source_directory}"
        )

    if chunk_size < 1:
        raise ValueError("chunk_size must be at least 1")

    if chunk_overlap < 0 or chunk_overlap >= chunk_size:
        raise ValueError("chunk_overlap must be between 0 and chunk_size - 1")

    if batch_size < 1:
        raise ValueError("batch_size must be at least 1")

    if max_workers < 1:
        raise ValueError("max_workers must be at least 1")

    source_files = sorted(source_directory.rglob("*.md"))
    total_files = len(source_files)
    all_chunks = []

    if total_files == 0:
        _display_progress("Chunking documents", 0, 0)

    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        for start in range(0, total_files, batch_size):
            batch = source_files[start : start + batch_size]
            batch_number = (start // batch_size) + 1
            batch_label = f"Chunking documents - Batch {batch_number}"
            completed_in_batch = 0
            _display_progress(batch_label, 0, len(batch))
            futures = {
                executor.submit(
                    _chunk_file,
                    file_path,
                    source_directory,
                    chunk_size,
                    chunk_overlap,
                )
                for file_path in batch
            }

            for future in as_completed(futures):
                file_chunks = future.result()
                all_chunks.extend(file_chunks)
                completed_in_batch += 1
                _display_progress(
                    batch_label,
                    completed_in_batch,
                    len(batch),
                )

    all_chunks.sort(key=lambda chunk: (chunk.source_file, chunk.chunk_index))
    return all_chunks


def _read_text_file(file_path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return file_path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue

    raise UnicodeError(f"Unable to decode: {file_path}")


def _csv_records(file_path: Path) -> list[tuple[str, str]]:
    raw_content = _read_text_file(file_path)

    try:
        dialect = csv.Sniffer().sniff(raw_content[:4096])
    except csv.Error:
        dialect = csv.excel

    reader = csv.DictReader(io.StringIO(raw_content), dialect=dialect)
    records = []

    for row_number, row in enumerate(reader, start=2):
        fields = [
            f"{column}: {value.strip()}"
            for column, value in row.items()
            if column and value and value.strip()
        ]

        if fields:
            records.append((f"Row {row_number}", "\n".join(fields)))

    return records


def _is_simple_json_value(value: object) -> bool:
    return not isinstance(value, (dict, list)) or (
        isinstance(value, list)
        and all(not isinstance(item, (dict, list)) for item in value)
    )


def _json_records(data: object) -> list[tuple[str, object]]:
    if isinstance(data, list):
        return [(f"$[{index}]", value) for index, value in enumerate(data)]

    if not isinstance(data, dict):
        return [("$", data)]

    simple_fields = {
        key: value for key, value in data.items() if _is_simple_json_value(value)
    }
    complex_fields = {
        key: value for key, value in data.items() if not _is_simple_json_value(value)
    }

    if not complex_fields:
        return [("$", data)]

    records = []

    for key, value in complex_fields.items():
        if isinstance(value, list):
            for index, item in enumerate(value):
                records.append(
                    (f"$.{key}[{index}]", {**simple_fields, key: item})
                )
        else:
            records.append((f"$.{key}", {**simple_fields, key: value}))

    return records


def _flatten_json(value: object, prefix: str = "") -> list[str]:
    if isinstance(value, dict):
        lines = []
        for key, nested_value in value.items():
            nested_prefix = f"{prefix}.{key}" if prefix else str(key)
            lines.extend(_flatten_json(nested_value, nested_prefix))
        return lines

    if isinstance(value, list):
        if all(not isinstance(item, (dict, list)) for item in value):
            return [f"{prefix}: {', '.join(map(str, value))}"]

        lines = []
        for index, nested_value in enumerate(value):
            lines.extend(_flatten_json(nested_value, f"{prefix}[{index}]"))
        return lines

    label = prefix or "value"
    return [f"{label}: {value}"]


def _load_json_records(file_path: Path) -> list[tuple[str, str]]:
    data = json.loads(_read_text_file(file_path))
    return [
        (json_path, "\n".join(_flatten_json(record)))
        for json_path, record in _json_records(data)
    ]


def _unique_headers(values: tuple[object, ...]) -> list[str]:
    headers = []
    occurrences: dict[str, int] = {}

    for column_number, value in enumerate(values, start=1):
        base_header = str(value).strip() if value is not None else ""
        base_header = base_header or f"Column {column_number}"
        occurrences[base_header] = occurrences.get(base_header, 0) + 1
        occurrence = occurrences[base_header]
        header = (
            base_header if occurrence == 1 else f"{base_header} ({occurrence})"
        )
        headers.append(header)

    return headers


def _load_xlsx_records(file_path: Path) -> list[tuple[str, str]]:
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    records = []

    try:
        for worksheet in workbook.worksheets:
            rows = enumerate(worksheet.iter_rows(values_only=True), start=1)
            header_row_number = None
            header_values = None

            for row_number, row in rows:
                if any(value is not None for value in row):
                    header_row_number = row_number
                    header_values = row
                    break

            if header_row_number is None or header_values is None:
                continue

            headers = _unique_headers(header_values)

            for row_index, row in rows:
                fields = [
                    f"{header}: {value}"
                    for header, value in zip(headers, row)
                    if value is not None and str(value).strip()
                ]

                if fields:
                    section = f"Sheet: {worksheet.title} > Row {row_index}"
                    records.append((section, "\n".join(fields)))
    finally:
        workbook.close()

    return records


def _chunk_structured_file(
    file_path: Path,
    source_directory: Path,
    chunk_size: int,
    chunk_overlap: int,
) -> list[DocumentChunk]:
    loaders = {
        ".csv": _csv_records,
        ".json": _load_json_records,
        ".xlsx": _load_xlsx_records,
    }
    extension = file_path.suffix.lower()
    records = loaders[extension](file_path)
    source_file = file_path.relative_to(source_directory).as_posix()
    chunks = []

    for section, record_content in records:
        for piece in _split_by_words(record_content, chunk_size, chunk_overlap):
            chunk_index = len(chunks)
            chunk_content = f"Record: {section}\n\n{piece}"
            chunk_id = sha256(
                f"{extension}:{source_file}:{chunk_index}:{chunk_content}".encode(
                    "utf-8"
                )
            ).hexdigest()[:16]
            chunks.append(
                DocumentChunk(
                    chunk_id=chunk_id,
                    content=chunk_content,
                    source_file=source_file,
                    source_type=extension.lstrip("."),
                    section=section,
                    chunk_index=chunk_index,
                )
            )

    return chunks


def chunk_structured_data(
    input_directory: str | Path,
    chunk_size: int,
    chunk_overlap: int,
    batch_size: int,
    max_workers: int,
) -> list[DocumentChunk]:
    """Chunk CSV, JSON, and XLSX records in parallel batches."""
    source_directory = Path(input_directory)

    if not source_directory.is_dir():
        raise NotADirectoryError(
            f"Structured data directory not found: {source_directory}"
        )

    if chunk_size < 1:
        raise ValueError("chunk_size must be at least 1")

    if chunk_overlap < 0 or chunk_overlap >= chunk_size:
        raise ValueError("chunk_overlap must be between 0 and chunk_size - 1")

    if batch_size < 1:
        raise ValueError("batch_size must be at least 1")

    if max_workers < 1:
        raise ValueError("max_workers must be at least 1")

    source_files = [
        path
        for path in sorted(source_directory.rglob("*"))
        if path.is_file() and path.suffix.lower() in STRUCTURED_EXTENSIONS
    ]
    total_files = len(source_files)
    all_chunks = []

    if total_files == 0:
        _display_progress("Chunking structured data", 0, 0)

    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        for start in range(0, total_files, batch_size):
            batch = source_files[start : start + batch_size]
            batch_number = (start // batch_size) + 1
            batch_label = f"Structured data - Batch {batch_number}"
            completed_in_batch = 0
            _display_progress(batch_label, 0, len(batch))
            futures = {
                executor.submit(
                    _chunk_structured_file,
                    file_path,
                    source_directory,
                    chunk_size,
                    chunk_overlap,
                )
                for file_path in batch
            }

            for future in as_completed(futures):
                file_chunks = future.result()
                all_chunks.extend(file_chunks)
                completed_in_batch += 1
                _display_progress(
                    batch_label,
                    completed_in_batch,
                    len(batch),
                )

    all_chunks.sort(
        key=lambda chunk: (chunk.source_type, chunk.source_file, chunk.chunk_index)
    )
    return all_chunks


def chunk_all_data(
    markdown_directory: str | Path,
    structured_directory: str | Path,
    chunk_size: int,
    chunk_overlap: int,
    batch_size: int,
    max_workers: int,
) -> list[DocumentChunk]:
    markdown_chunks = chunk_documents(
        input_directory=markdown_directory,
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        batch_size=batch_size,
        max_workers=max_workers,
    )
    structured_chunks = chunk_structured_data(
        input_directory=structured_directory,
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        batch_size=batch_size,
        max_workers=max_workers,
    )
    return markdown_chunks + structured_chunks
