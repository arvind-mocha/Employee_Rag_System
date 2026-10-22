from pathlib import Path
import csv
import io
import json

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from openpyxl import load_workbook
from pypdf import PdfReader


def _read_text_file(file_path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return file_path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue

    raise UnicodeError(f"Unable to decode: {file_path}")


def _extract_pdf(file_path: Path) -> str:
    reader = PdfReader(file_path)
    pages = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        pages.append(f"[Page {page_number}]\n{text.strip()}")

    return "\n\n".join(pages)


def _extract_docx(file_path: Path) -> str:
    document = Document(file_path)
    content = []

    for element in document.element.body.iterchildren():
        if isinstance(element, CT_P):
            paragraph = Paragraph(element, document)
            text = paragraph.text.strip()

            if text:
                content.append(text)

        elif isinstance(element, CT_Tbl):
            table = Table(element, document)

            for row in table.rows:
                values = [
                    " ".join(cell.text.split())
                    for cell in row.cells
                ]
                content.append("\t".join(values))

    return "\n".join(content)


def _extract_xlsx(file_path: Path) -> str:
    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=False,
    )

    content = []

    try:
        for worksheet in workbook.worksheets:
            content.append(f"[Sheet: {worksheet.title}]")

            for row in worksheet.iter_rows(values_only=True):
                values = [
                    "" if value is None else str(value)
                    for value in row
                ]

                if any(values):
                    content.append("\t".join(values))

            content.append("")
    finally:
        workbook.close()

    return "\n".join(content).strip()


def _extract_csv(file_path: Path) -> str:
    raw_text = _read_text_file(file_path)

    try:
        dialect = csv.Sniffer().sniff(raw_text[:4096])
    except csv.Error:
        dialect = csv.excel

    reader = csv.reader(io.StringIO(raw_text), dialect)
    return "\n".join(
        "\t".join(cell.strip() for cell in row)
        for row in reader
    )


def _extract_json(file_path: Path) -> str:
    raw_text = _read_text_file(file_path)
    data = json.loads(raw_text)

    return json.dumps(
        data,
        ensure_ascii=False,
        indent=2,
    )

def extract_content(file_path: str | Path) -> str:
    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError(f"File not found: {path}")

    extension = path.suffix.lower()

    extractors = {
        ".pdf": _extract_pdf,
        ".docx": _extract_docx,
        ".xlsx": _extract_xlsx,
        ".csv": _extract_csv,
        ".json": _extract_json,
        ".txt": _read_text_file,
        ".md": _read_text_file,
        ".markdown": _read_text_file,
    }

    extractor = extractors.get(extension)

    if extractor is None:
        raise ValueError(f"Unsupported file type: {extension}")

    return extractor(path)